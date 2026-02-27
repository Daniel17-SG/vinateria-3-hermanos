from django.contrib import admin
from django.utils.html import format_html
from django.http import HttpResponse
from .models import Categoria, Producto, PerfilCliente, CarritoItem, Venta, DetalleVenta, MensajeContacto
from . import admin_custom
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime

# ¡ELIMINAMOS LA LÍNEA admin.site = admin_custom.admin_site!

def exportar_ventas_excel(modeladmin, request, queryset):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ventas"
    
    headers = ['ID', 'Fecha', 'Usuario', 'Total', 'Estatus', 'Dirección', 'Teléfono']
    ws.append(headers)
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="8B0000", end_color="8B0000", fill_type="solid")
    
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    
    for venta in queryset:
        ws.append([
            venta.id,
            venta.fecha_venta.strftime('%Y-%m-%d %H:%M'),
            venta.usuario.username,
            float(venta.total),
            venta.get_estatus_display(),
            venta.direccion_envio,
            venta.telefono_contacto
        ])
    
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        ws.column_dimensions[column].width = max_length + 2
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="ventas_{datetime.now().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response

exportar_ventas_excel.short_description = "Exportar ventas a Excel"

class DateRangeFilter(admin.SimpleListFilter):
    title = 'rango de fechas'
    parameter_name = 'fecha_rango'
    
    def lookups(self, request, model_admin):
        return [
            ('hoy', 'Hoy'),
            ('semana', 'Esta semana'),
            ('mes', 'Este mes'),
        ]
    
    def queryset(self, request, queryset):
        today = datetime.now().date()
        if self.value() == 'hoy':
            return queryset.filter(fecha_venta__date=today)
        elif self.value() == 'semana':
            from datetime import timedelta
            return queryset.filter(fecha_venta__date__gte=today - timedelta(days=7))
        elif self.value() == 'mes':
            return queryset.filter(fecha_venta__date__gte=today.replace(day=1))

# NOTA CÓMO AHORA EL DECORADOR APUNTA DIRECTO A TU ADMIN_CUSTOM
@admin.register(MensajeContacto, site=admin_custom.admin_site)
class MensajeContactoAdmin(admin.ModelAdmin):
    list_display = ('nombre', 'email', 'mensaje_corto', 'fecha_envio', 'leido')
    list_filter = ('leido', 'fecha_envio')
    search_fields = ('nombre', 'email', 'mensaje')
    list_editable = ('leido',)
    readonly_fields = ('fecha_envio',)
    date_hierarchy = 'fecha_envio'
    
    def mensaje_corto(self, obj):
        return obj.mensaje[:50] + '...' if len(obj.mensaje) > 50 else obj.mensaje
    mensaje_corto.short_description = 'Mensaje'

@admin.register(Categoria, site=admin_custom.admin_site)
class CategoriaAdmin(admin.ModelAdmin):
    list_display = ('nombre', 'slug', 'activo', 'productos_count')
    list_filter = ('activo',)
    search_fields = ('nombre', 'descripcion')
    prepopulated_fields = {'slug': ('nombre',)}
    list_editable = ('activo',)
    
    def productos_count(self, obj):
        return obj.productos.count()
    productos_count.short_description = 'Productos'

@admin.register(Producto, site=admin_custom.admin_site)
class ProductoAdmin(admin.ModelAdmin):
    list_display = ('thumbnail', 'nombre', 'categoria', 'precio', 'stock', 'stock_badge', 'activo', 'fecha_creacion')
    list_filter = ('categoria', 'activo', 'fecha_creacion')
    search_fields = ('nombre', 'descripcion')
    list_editable = ('precio', 'stock', 'activo')
    prepopulated_fields = {'slug': ('nombre',)}
    readonly_fields = ('fecha_creacion', 'fecha_actualizacion')
    list_per_page = 25
    date_hierarchy = 'fecha_creacion'
    
    fieldsets = (
        ('Información General', {
            'fields': ('nombre', 'slug', 'categoria', 'descripcion')
        }),
        ('Precios y Stock', {
            'fields': ('precio', 'stock', 'activo')
        }),
        ('Imagen', {
            'fields': ('imagen',)
        }),
        ('Fechas', {
            'fields': ('fecha_creacion', 'fecha_actualizacion'),
            'classes': ('collapse',)
        }),
    )
    
    def thumbnail(self, obj):
        if obj.imagen:
            return format_html('<img src="{}" style="width: 50px; height: 50px; object-fit: cover; border-radius: 4px;">', obj.imagen.url)
        return format_html('<span style="color: #999;">Sin imagen</span>')
    thumbnail.short_description = 'Imagen'
    
    def stock_badge(self, obj):
        if obj.stock == 0:
            return format_html('<span style="background-color: #dc3545; color: white; padding: 3px 8px; border-radius: 4px; font-size: 12px;">AGOTADO</span>')
        elif obj.stock < 10:
            return format_html('<span style="background-color: #ffc107; color: #333; padding: 3px 8px; border-radius: 4px; font-size: 12px;">BAJO ({})</span>', obj.stock)
        return obj.stock
    stock_badge.short_description = 'Stock'

@admin.register(PerfilCliente, site=admin_custom.admin_site)
class PerfilClienteAdmin(admin.ModelAdmin):
    list_display = ('user', 'email', 'telefono', 'ciudad', 'cp', 'fecha_registro')
    search_fields = ('user__username', 'user__email', 'telefono', 'ciudad')
    list_filter = ('ciudad',)
    readonly_fields = ('fecha_registro',)
    
    def email(self, obj):
        return obj.user.email
    email.short_description = 'Email'

@admin.register(CarritoItem, site=admin_custom.admin_site)
class CarritoItemAdmin(admin.ModelAdmin):
    list_display = ('usuario', 'producto', 'cantidad', 'fecha_agregado')
    list_filter = ('fecha_agregado',)
    search_fields = ('usuario__username', 'producto__nombre')
    date_hierarchy = 'fecha_agregado'

@admin.register(Venta, site=admin_custom.admin_site)
class VentaAdmin(admin.ModelAdmin):
    list_display = ('id', 'usuario', 'total', 'estatus', 'estatus_badge', 'fecha_venta')
    list_filter = ('estatus', 'fecha_venta', DateRangeFilter)
    search_fields = ('usuario__username', 'id', 'telefono_contacto')
    list_editable = ('estatus',)
    readonly_fields = ('fecha_venta',)
    date_hierarchy = 'fecha_venta'
    list_per_page = 20
    
    ESTATUS_COLORS = {
        'pendiente': 'orange',
        'pagado': 'blue',
        'enviado': 'purple',
        'entregado': 'green',
        'cancelado': 'red',
    }
    
    def estatus_badge(self, obj):
        color = self.ESTATUS_COLORS.get(obj.estatus, 'gray')
        return format_html(
            '<span style="background-color: {}; color: white; padding: 3px 8px; border-radius: 4px; font-size: 12px;">{}</span>',
            color, obj.get_estatus_display()
        )
    estatus_badge.short_description = 'Estatus'
    
    actions = ['marcar_como_pagado', 'marcar_como_enviado', 'marcar_como_entregado', 'exportar_ventas_excel']
    
    def marcar_como_pagado(self, request, queryset):
        queryset.update(estatus='pagado')
    marcar_como_pagado.short_description = 'Marcar como Pagado'
    
    def marcar_como_enviado(self, request, queryset):
        queryset.update(estatus='enviado')
    marcar_como_enviado.short_description = 'Marcar como Enviado'
    
    def marcar_como_entregado(self, request, queryset):
        queryset.update(estatus='entregado')
    marcar_como_entregado.short_description = 'Marcar como Entregado'

@admin.register(DetalleVenta, site=admin_custom.admin_site)
class DetalleVentaAdmin(admin.ModelAdmin):
    list_display = ('venta', 'producto', 'cantidad', 'precio_unitario', 'subtotal_display')
    search_fields = ('venta__id', 'producto__nombre')
    readonly_fields = ('venta', 'producto', 'cantidad', 'precio_unitario')
    
    def subtotal_display(self, obj):
        return f"${obj.subtotal}"
    subtotal_display.short_description = 'Subtotal'